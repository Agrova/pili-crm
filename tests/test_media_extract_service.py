"""ADR-014 Task 5: tests for analysis/media_extract/service.py.

Phase A covers the selector, routing, office/placeholder branch, and the
idempotent writer. Phase B (vision) extends this same module.
"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import AsyncMock, patch

import docx
import openpyxl
import pytest
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from analysis.media_extract import service
from analysis.media_extract.service import (
    ExtractionResult,
    ExtractorKind,
    PendingMediaMessage,
    decide_extractor,
    extract_office_or_placeholder,
    save_extraction,
    select_pending_messages,
)
from analysis.media_extract.vision import VisionAPIError, VisionExtractionResult, VisionImageError

EXTRACTOR_VERSION = "v1.0-test"
SEED_PHONE = "+77471057849"


# ── DB fixture helpers ─────────────────────────────────────────────────────


async def _seed_account_id(session: AsyncSession) -> int:
    aid = (
        await session.execute(
            text(
                "SELECT id FROM communications_telegram_account "
                "WHERE phone_number = :phone"
            ),
            {"phone": SEED_PHONE},
        )
    ).scalar()
    assert aid is not None, "ADR-012 seed account missing"
    return int(aid)


async def _seed_chat(session: AsyncSession, account_id: int, tag: str) -> int:
    cid = (
        await session.execute(
            text(
                "INSERT INTO communications_telegram_chat "
                "(owner_account_id, telegram_chat_id, chat_type, title) "
                "VALUES (:aid, :tg, 'personal_chat', :title) RETURNING id"
            ),
            {
                "aid": account_id,
                "tg": f"tg-svc-{tag}-{datetime.now(tz=UTC).timestamp()}",
                "title": f"svc test {tag}",
            },
        )
    ).scalar_one()
    await session.flush()
    return int(cid)


async def _seed_message(
    session: AsyncSession,
    chat_id: int,
    *,
    media_type: str,
    mime_type: str | None = None,
    file_name: str | None = None,
    relative_path: str | None = None,
    file_size_bytes: int | None = None,
    tag: str = "msg",
) -> int:
    mid = (
        await session.execute(
            text(
                "INSERT INTO communications_telegram_message "
                "(chat_id, telegram_message_id, sent_at) "
                "VALUES (:cid, :tg, NOW()) RETURNING id"
            ),
            {
                "cid": chat_id,
                "tg": f"svc-msg-{tag}-{datetime.now(tz=UTC).timestamp()}",
            },
        )
    ).scalar_one()
    await session.execute(
        text(
            "INSERT INTO communications_telegram_message_media "
            "(message_id, media_type, mime_type, file_name, "
            " relative_path, file_size_bytes) "
            "VALUES (:mid, :media_type, :mime_type, :file_name, "
            "        :relative_path, :file_size_bytes)"
        ),
        {
            "mid": mid,
            "media_type": media_type,
            "mime_type": mime_type,
            "file_name": file_name,
            "relative_path": relative_path,
            "file_size_bytes": file_size_bytes,
        },
    )
    await session.flush()
    return int(mid)


# ── 1. select_pending_messages ─────────────────────────────────────────────


async def test_select_pending_messages_filters_by_chat_id(
    db_session: AsyncSession,
) -> None:
    aid = await _seed_account_id(db_session)
    chat_a = await _seed_chat(db_session, aid, tag="A")
    chat_b = await _seed_chat(db_session, aid, tag="B")
    msg_a = await _seed_message(
        db_session, chat_a, media_type="photo", relative_path="chats/a/p.jpg",
        tag="a-photo",
    )
    msg_b = await _seed_message(
        db_session, chat_b, media_type="photo", relative_path="chats/b/p.jpg",
        tag="b-photo",
    )

    found = await select_pending_messages(
        db_session,
        chat_id=chat_a,
        extractor_version=EXTRACTOR_VERSION,
        batch_size=100,
    )
    ids = {m.message_id for m in found}
    assert msg_a in ids
    assert msg_b not in ids


async def test_select_pending_messages_skips_existing(
    db_session: AsyncSession,
) -> None:
    aid = await _seed_account_id(db_session)
    chat = await _seed_chat(db_session, aid, tag="skip")
    msg = await _seed_message(
        db_session, chat, media_type="photo", relative_path="x/p.jpg",
        tag="skip",
    )

    await db_session.execute(
        text(
            "INSERT INTO communications_telegram_message_media_extraction "
            "(message_id, extracted_text, extraction_method, extractor_version) "
            "VALUES (:mid, 'done', 'placeholder', :ver)"
        ),
        {"mid": msg, "ver": EXTRACTOR_VERSION},
    )
    await db_session.flush()

    found = await select_pending_messages(
        db_session,
        chat_id=chat,
        extractor_version=EXTRACTOR_VERSION,
        batch_size=100,
        skip_existing=True,
    )
    assert msg not in {m.message_id for m in found}


async def test_select_pending_messages_returns_existing_when_skip_false(
    db_session: AsyncSession,
) -> None:
    aid = await _seed_account_id(db_session)
    chat = await _seed_chat(db_session, aid, tag="regen")
    msg = await _seed_message(
        db_session, chat, media_type="photo", relative_path="x/p.jpg",
        tag="regen",
    )
    await db_session.execute(
        text(
            "INSERT INTO communications_telegram_message_media_extraction "
            "(message_id, extracted_text, extraction_method, extractor_version) "
            "VALUES (:mid, 'done', 'placeholder', :ver)"
        ),
        {"mid": msg, "ver": EXTRACTOR_VERSION},
    )
    await db_session.flush()

    found = await select_pending_messages(
        db_session,
        chat_id=chat,
        extractor_version=EXTRACTOR_VERSION,
        batch_size=100,
        skip_existing=False,
    )
    assert msg in {m.message_id for m in found}


async def test_select_pending_messages_resolves_phone_number_via_join(
    db_session: AsyncSession,
) -> None:
    aid = await _seed_account_id(db_session)
    chat = await _seed_chat(db_session, aid, tag="phone")
    msg = await _seed_message(
        db_session, chat, media_type="photo",
        relative_path="chats/c/p.jpg", tag="phone",
    )

    found = await select_pending_messages(
        db_session,
        chat_id=chat,
        extractor_version=EXTRACTOR_VERSION,
        batch_size=10,
    )
    target = next(m for m in found if m.message_id == msg)
    assert target.phone_number == SEED_PHONE
    assert target.relative_path == "chats/c/p.jpg"
    assert target.absolute_path(Path("/tmp/exports")) == Path(
        f"/tmp/exports/{SEED_PHONE}/chats/c/p.jpg"
    )


# ── 2. decide_extractor ────────────────────────────────────────────────────


def _msg(**overrides: object) -> PendingMediaMessage:
    base = dict(
        message_id=1,
        media_type="photo",
        mime_type=None,
        file_name=None,
        relative_path="some/path",
        file_size_bytes=None,
        phone_number="+70000000000",
    )
    base.update(overrides)
    return PendingMediaMessage(**base)  # type: ignore[arg-type]


def test_decide_extractor_photo() -> None:
    assert decide_extractor(_msg(media_type="photo")) is ExtractorKind.VISION


def test_decide_extractor_image_file() -> None:
    assert (
        decide_extractor(
            _msg(media_type="file", mime_type="image/jpeg", file_name="x.jpg")
        )
        is ExtractorKind.VISION
    )


def test_decide_extractor_xlsx_by_mime() -> None:
    assert (
        decide_extractor(
            _msg(
                media_type="file",
                mime_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                file_name="order.bin",
            )
        )
        is ExtractorKind.XLSX
    )


def test_decide_extractor_xlsx_by_extension() -> None:
    assert (
        decide_extractor(
            _msg(media_type="file", mime_type=None, file_name="ORDER.XLSX")
        )
        is ExtractorKind.XLSX
    )


def test_decide_extractor_docx_by_mime() -> None:
    assert (
        decide_extractor(
            _msg(
                media_type="file",
                mime_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                file_name="doc.bin",
            )
        )
        is ExtractorKind.DOCX
    )


def test_decide_extractor_video_returns_placeholder() -> None:
    assert (
        decide_extractor(_msg(media_type="video_file", mime_type="video/mp4"))
        is ExtractorKind.PLACEHOLDER
    )


def test_decide_extractor_unknown_returns_placeholder() -> None:
    assert (
        decide_extractor(
            _msg(media_type="file", mime_type="application/zip", file_name="x.zip")
        )
        is ExtractorKind.PLACEHOLDER
    )


def test_decide_extractor_no_relative_path_returns_placeholder() -> None:
    assert (
        decide_extractor(_msg(media_type="photo", relative_path=None))
        is ExtractorKind.PLACEHOLDER
    )


# ── 3. extract_office_or_placeholder ───────────────────────────────────────


async def test_extract_office_xlsx_success(tmp_path: Path) -> None:
    exports_root = tmp_path / "exports"
    file_dir = exports_root / SEED_PHONE / "chats" / "c"
    file_dir.mkdir(parents=True)
    xlsx_path = file_dir / "order.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["A", "B"])
    ws.append(["1", "2"])
    wb.save(xlsx_path)

    msg = _msg(
        message_id=42,
        media_type="file",
        mime_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        file_name="order.xlsx",
        relative_path="chats/c/order.xlsx",
        phone_number=SEED_PHONE,
    )
    result = await extract_office_or_placeholder(
        msg, ExtractorKind.XLSX, exports_root
    )
    assert result.message_id == 42
    assert result.extraction_method == "xlsx_openpyxl"
    assert "order.xlsx" in result.extracted_text
    assert "Sheet1" in result.extracted_text


async def test_extract_office_xlsx_corrupted_returns_placeholder(
    tmp_path: Path,
) -> None:
    exports_root = tmp_path / "exports"
    file_dir = exports_root / SEED_PHONE / "x"
    file_dir.mkdir(parents=True)
    xlsx_path = file_dir / "broken.xlsx"
    xlsx_path.write_bytes(b"not actually xlsx")

    msg = _msg(
        message_id=99,
        media_type="file",
        mime_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        file_name="broken.xlsx",
        relative_path="x/broken.xlsx",
        file_size_bytes=18,
        phone_number=SEED_PHONE,
    )
    result = await extract_office_or_placeholder(
        msg, ExtractorKind.XLSX, exports_root
    )
    assert result.extraction_method == "placeholder"
    assert "broken.xlsx" in result.extracted_text
    assert "parse error" in result.extracted_text


async def test_extract_office_file_not_on_disk_returns_placeholder(
    tmp_path: Path,
) -> None:
    exports_root = tmp_path / "exports"
    msg = _msg(
        message_id=7,
        media_type="file",
        mime_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml."
            "document"
        ),
        file_name="ghost.docx",
        relative_path="chats/g/ghost.docx",
        file_size_bytes=0,
        phone_number=SEED_PHONE,
    )
    result = await extract_office_or_placeholder(
        msg, ExtractorKind.DOCX, exports_root
    )
    assert result.extraction_method == "placeholder"
    assert "ghost.docx" in result.extracted_text
    assert "file not found on disk" in result.extracted_text


async def test_extract_office_docx_success(tmp_path: Path) -> None:
    exports_root = tmp_path / "exports"
    file_dir = exports_root / SEED_PHONE / "chats" / "d"
    file_dir.mkdir(parents=True)
    docx_path = file_dir / "letter.docx"
    document = docx.Document()
    document.add_paragraph("Hello world")
    document.save(docx_path)

    msg = _msg(
        message_id=11,
        media_type="file",
        mime_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml."
            "document"
        ),
        file_name="letter.docx",
        relative_path="chats/d/letter.docx",
        phone_number=SEED_PHONE,
    )
    result = await extract_office_or_placeholder(
        msg, ExtractorKind.DOCX, exports_root
    )
    assert result.extraction_method == "docx_python_docx"
    assert "Hello world" in result.extracted_text


async def test_extract_office_placeholder_format_for_unknown_type() -> None:
    msg = _msg(
        message_id=1,
        media_type="file",
        mime_type="application/pdf",
        file_name="report.pdf",
        relative_path="chats/x/report.pdf",
        file_size_bytes=2048,
    )
    result = await extract_office_or_placeholder(
        msg, ExtractorKind.PLACEHOLDER, Path("/nope")
    )
    assert result.extraction_method == "placeholder"
    assert (
        result.extracted_text
        == "[file: report.pdf, type: application/pdf, size: 2048 bytes]"
    )


async def test_extract_office_placeholder_no_size_no_path() -> None:
    msg = _msg(
        message_id=2,
        media_type="file",
        mime_type="application/zip",
        file_name="x.zip",
        relative_path=None,
        file_size_bytes=None,
    )
    result = await extract_office_or_placeholder(
        msg, ExtractorKind.PLACEHOLDER, Path("/nope")
    )
    assert "size: unknown" in result.extracted_text
    assert "(file not exported)" in result.extracted_text


# ── 4. save_extraction ─────────────────────────────────────────────────────


async def test_save_extraction_inserts_new(db_session: AsyncSession) -> None:
    aid = await _seed_account_id(db_session)
    chat = await _seed_chat(db_session, aid, tag="ins")
    msg = await _seed_message(
        db_session, chat, media_type="photo", relative_path="x/p.jpg", tag="ins",
    )

    inserted = await save_extraction(
        db_session,
        ExtractionResult(
            message_id=msg,
            extracted_text="[Изображение]\nОписание: x\nТекст на изображении: y",
            extraction_method="vision_qwen3-vl-30b-a3b",
        ),
        EXTRACTOR_VERSION,
    )
    assert inserted is True

    row = (
        await db_session.execute(
            text(
                "SELECT extraction_method, extractor_version "
                "FROM communications_telegram_message_media_extraction "
                "WHERE message_id = :mid"
            ),
            {"mid": msg},
        )
    ).one()
    assert row.extraction_method == "vision_qwen3-vl-30b-a3b"
    assert row.extractor_version == EXTRACTOR_VERSION


async def test_save_extraction_idempotent_on_conflict(
    db_session: AsyncSession,
) -> None:
    aid = await _seed_account_id(db_session)
    chat = await _seed_chat(db_session, aid, tag="conf")
    msg = await _seed_message(
        db_session, chat, media_type="photo", relative_path="x/p.jpg", tag="conf",
    )
    first = await save_extraction(
        db_session,
        ExtractionResult(
            message_id=msg, extracted_text="first",
            extraction_method="placeholder",
        ),
        EXTRACTOR_VERSION,
    )
    second = await save_extraction(
        db_session,
        ExtractionResult(
            message_id=msg, extracted_text="second",
            extraction_method="placeholder",
        ),
        EXTRACTOR_VERSION,
    )
    assert first is True
    assert second is False

    stored = (
        await db_session.execute(
            text(
                "SELECT extracted_text "
                "FROM communications_telegram_message_media_extraction "
                "WHERE message_id = :mid"
            ),
            {"mid": msg},
        )
    ).scalar()
    assert stored == "first"


async def test_save_extraction_regenerate_deletes_old(
    db_session: AsyncSession,
) -> None:
    aid = await _seed_account_id(db_session)
    chat = await _seed_chat(db_session, aid, tag="regen")
    msg = await _seed_message(
        db_session, chat, media_type="photo", relative_path="x/p.jpg", tag="regen",
    )
    await save_extraction(
        db_session,
        ExtractionResult(
            message_id=msg, extracted_text="old",
            extraction_method="placeholder",
        ),
        "v0.9",
    )
    await save_extraction(
        db_session,
        ExtractionResult(
            message_id=msg, extracted_text="new",
            extraction_method="vision_qwen3-vl-30b-a3b",
        ),
        EXTRACTOR_VERSION,
        regenerate=True,
    )

    row = (
        await db_session.execute(
            text(
                "SELECT extracted_text, extraction_method, extractor_version "
                "FROM communications_telegram_message_media_extraction "
                "WHERE message_id = :mid"
            ),
            {"mid": msg},
        )
    ).one()
    assert row.extracted_text == "new"
    assert row.extraction_method == "vision_qwen3-vl-30b-a3b"
    assert row.extractor_version == EXTRACTOR_VERSION


# ──────────────────────────────────────────────────────────────────────────
# Phase B — vision wrapper
# ──────────────────────────────────────────────────────────────────────────


def _vision_msg(message_id: int) -> PendingMediaMessage:
    return PendingMediaMessage(
        message_id=message_id,
        media_type="photo",
        mime_type="image/jpeg",
        file_name="p.jpg",
        relative_path="chats/v/p.jpg",
        file_size_bytes=1024,
        phone_number=SEED_PHONE,
    )


async def test_extract_image_success_returns_vision_result(tmp_path: Path) -> None:
    exports_root = tmp_path / "exports"
    file_dir = exports_root / SEED_PHONE / "chats" / "v"
    file_dir.mkdir(parents=True)
    (file_dir / "p.jpg").write_bytes(b"fake-bytes")

    fake_text = "[Изображение]\nОписание: x\nТекст на изображении: y"
    fake_result = VisionExtractionResult(text=fake_text, extraction_method="vision")
    with patch.object(
        service, "_vision_extract_image", new=AsyncMock(return_value=fake_result)
    ):
        result = await service.extract_image_or_fail(
            _vision_msg(33),
            exports_root,
            model_id="qwen/qwen3-vl-30b",
            endpoint="http://localhost:1234/v1",
        )
    assert result.message_id == 33
    assert result.extracted_text == fake_text
    assert result.extraction_method == "vision_qwen3-vl-30b-a3b"


async def test_extract_image_file_not_found_returns_placeholder(
    tmp_path: Path,
) -> None:
    exports_root = tmp_path / "exports"
    result = await service.extract_image_or_fail(
        _vision_msg(44),
        exports_root,
        model_id="qwen/qwen3-vl-30b",
        endpoint="http://localhost:1234/v1",
    )
    assert result.extraction_method == "placeholder"
    assert "file not found on disk" in result.extracted_text


async def test_extract_image_VisionImageError_returns_placeholder(
    tmp_path: Path,
) -> None:
    exports_root = tmp_path / "exports"
    file_dir = exports_root / SEED_PHONE / "chats" / "v"
    file_dir.mkdir(parents=True)
    (file_dir / "p.jpg").write_bytes(b"not-an-image")

    with patch.object(
        service,
        "_vision_extract_image",
        new=AsyncMock(side_effect=VisionImageError("decode failed")),
    ):
        result = await service.extract_image_or_fail(
            _vision_msg(55),
            exports_root,
            model_id="qwen/qwen3-vl-30b",
            endpoint="http://localhost:1234/v1",
        )
    assert result.extraction_method == "placeholder"
    assert "parse error" in result.extracted_text


async def test_extract_image_VisionAPIError_propagates(tmp_path: Path) -> None:
    exports_root = tmp_path / "exports"
    file_dir = exports_root / SEED_PHONE / "chats" / "v"
    file_dir.mkdir(parents=True)
    (file_dir / "p.jpg").write_bytes(b"fake")

    with (
        patch.object(
            service,
            "_vision_extract_image",
            new=AsyncMock(side_effect=VisionAPIError("OOM")),
        ),
        pytest.raises(VisionAPIError),
    ):
        await service.extract_image_or_fail(
            _vision_msg(66),
            exports_root,
            model_id="qwen/qwen3-vl-30b",
            endpoint="http://localhost:1234/v1",
        )


def test_derive_extraction_method_primary() -> None:
    assert (
        service.derive_extraction_method_from_model(
            "mlx-community/Qwen3-VL-30B-A3B-Instruct-4bit"
        )
        == "vision_qwen3-vl-30b-a3b"
    )


def test_derive_extraction_method_fallback() -> None:
    assert (
        service.derive_extraction_method_from_model(
            "mlx-community/Qwen3-VL-8B-Instruct-MLX-4bit"
        )
        == "vision_qwen3-vl-8b"
    )


def test_derive_extraction_method_unknown() -> None:
    out = service.derive_extraction_method_from_model(
        "openai/Custom_Model_v2.5-FP16"
    )
    assert out == "vision_custom-model-v2-5-fp16"


